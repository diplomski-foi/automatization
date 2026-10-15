"""Input loading: data tables (CSV / XLSX / JSON), URL lists and login accounts."""
import csv
import json
import re
from pathlib import Path


def _clean_row(row: dict) -> dict:
    """Strip keys/values and drop empty cells so they are not filled."""
    out = {}
    for k, v in row.items():
        if k is None:
            continue
        key = str(k).strip()
        if not key or v is None:
            continue
        if hasattr(v, "isoformat"):  # dates from Excel
            v = v.date().isoformat() if hasattr(v, "date") else v.isoformat()
        if isinstance(v, bool):
            v = "true" if v else "false"
        v = str(v).strip()
        if v != "":
            out[key] = v
    return out


def load_table(path) -> list[dict]:
    """Load rows from .csv / .tsv / .xlsx / .json (object or list of objects)."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
        rows = data if isinstance(data, list) else [data]
        return rows
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(it, [])]
        return [r for r in (_clean_row(dict(zip(header, row))) for row in it) if r]
    text = path.read_text(encoding="utf-8-sig")
    try:
        dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(text.splitlines(), dialect=dialect)
    return [r for r in (_clean_row(row) for row in reader) if r]


def parse_urls(text: str) -> list[str]:
    """One URL per line; blank lines and '#' comments are ignored."""
    urls = []
    for line in text.splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        urls.append(line)
    return urls


def load_urls(path) -> list[str]:
    """URLs from a .txt (one per line) or a .csv/.xlsx with a 'url' column."""
    path = Path(path)
    if path.suffix.lower() in (".csv", ".xlsx", ".tsv"):
        rows = load_table(path)
        return [r["url"] for r in rows if r.get("url")]
    return parse_urls(path.read_text(encoding="utf-8-sig"))


def load_logins(path) -> list[dict]:
    """Accounts to log in with.

    Each row: any credential columns (username, password, ...) plus an optional
    `login_url`. Returns [{"label", "login_url", "creds"}].
    """
    accounts = []
    for row in load_table(path):
        row = dict(row)
        login_url = row.pop("login_url", None) or row.pop("url", None)
        label = row.get("username") or row.get("email") or row.get("user") or "account"
        accounts.append({"label": label, "login_url": login_url, "creds": row})
    return accounts


def to_file_uri_if_path(value: str) -> str:
    """Allow plain local paths in URL lists."""
    if re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*://", value):
        return value
    p = Path(value).expanduser()
    return p.resolve().as_uri() if p.exists() else value
